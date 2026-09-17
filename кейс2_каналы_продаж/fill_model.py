# -*- coding: utf-8 -*-
"""Заполняет модель «Кейс 2. Каналы продаж» данными проекта (футболки с принтом).
Годовые цифры = месячный срез из шаблона Адамчика × 12."""
import openpyxl
from openpyxl.styles import Font

SRC = "Кейс2_модель_каналов_ШАБЛОН.xlsx"
DST = "Кейс2_Каналы_продаж_ФУТБОЛКИ.xlsx"

# канал, расходы, показы, переходы, интерес(корзина), заказы, выручка, валовая маржа канала
DATA = [
    ("Сайт (SEO и контент)",   60_000,   312_000,  3_120,   624, 156,   528_840, 2171/3390),
    ("Яндекс Директ",         144_000, 1_440_000, 14_400, 1_152, 144,   488_160, 2171/3390),
    ("LaModa",                144_000,   720_000, 21_600, 2_160, 360, 1_220_400, 1202/3390),
    ("Блогеры (интеграции)",  180_000,   960_000, 19_200, 1_920, 240,   732_240, 1839/3051),
    ("ВК Канал (Таргет ВК)",  120_000, 1_200_000,  8_400,   240,  84,   284_760, 1763/3390),
]
# Ozon и WB исключены: продаж на этих площадках не планируется

wb = openpyxl.load_workbook(SRC)
d = wb["Данные"]

d["A1"] = "Кейс 2. Каналы продаж — данные проекта «Спортивные футболки с принтом», год"
d["A2"] = ("Факт за 12 месяцев (месячный срез таблицы × 12). Юнит-экономика согласована с Кейсом 1: цена 3 390 ₽, "
           "себестоимость 851 ₽ (футболка 780 + принт 46 + упаковка 25), у блогеров промокод −10%.")

d["H4"] = "Валовая маржа канала"
d["H4"].font = d["G4"].font.copy()
d["H4"].fill = d["G4"].fill.copy()
d["H4"].border = d["G4"].border.copy()
d["H4"].alignment = d["G4"].alignment.copy()
d["I4"] = "доля выручки после себестоимости и комиссии канала"
d["I4"].font = Font(italic=True, size=9, color="7F7F7F")

for i, (name, rasx, pok, per, inter, zak, vyr, marg) in enumerate(DATA):
    r = 5 + i
    d.cell(row=r, column=1, value=name)
    d.cell(row=r, column=2, value=rasx)
    d.cell(row=r, column=3, value=pok)
    d.cell(row=r, column=4, value=per)
    d.cell(row=r, column=5, value=inter)
    d.cell(row=r, column=6, value=zak)
    d.cell(row=r, column=7, value=vyr)
    c = d.cell(row=r, column=8, value=round(marg, 4))
    c.number_format = "0.0%"
    c.border = d.cell(row=r, column=7).border.copy()
    c.fill = d.cell(row=r, column=7).fill.copy()
d["H12"] = "=IF(G12=0,0,SUMPRODUCT(G5:G11,H5:H11)/G12)"
d["H12"].number_format = "0.0%"
d["H12"].font = d["G12"].font.copy()

# средняя маржа по компании — параметр ROAS безубыточности
d["B16"] = "=H12"
d["C16"] = "средневзвешенная доля выручки после себестоимости и комиссий каналов"

# валовая прибыль и вердикт — по марже конкретного канала, а не по средней:
# у маркетплейсов комиссия 20–35%, усреднение завысило бы их ROMI
a = wb["Анализ"]
a["A2"] = ("ROAS = выручка / расходы. ROMI учитывает маржу канала: сколько валовой прибыли "
           "сверх вложенного приносит рубль. У каждого канала своя маржа — комиссии МП разные.")
for r in range(5, 12):
    a[f"O{r}"] = f"=C{r}*Данные!$H{r}"
    a[f"T{r}"] = (f'=IF(Данные!$H{r}=0,"",IF(M{r}<1/Данные!$H{r},"Убыточен — сократить",'
                  f'IF(M{r}>=$M$12,"Лидер — наращивать","Ниже среднего — оптимизировать")))')
a["A15"] = "ROAS безубыточности (1 / средняя маржа)"
a["C15"] = "по компании; у каждого канала свой порог — он зашит в колонку «Оценка канала»"


# в модели 7 слотов под каналы — лишние вычищаем, иначе в рейтинге и бюджете висят нули
free = range(5 + len(DATA), 12)                       # строки Данные/Анализ
for r in free:
    for col in range(1, 9):
        d.cell(row=r, column=col).value = None
        a.cell(row=r, column=col).value = None
    for col in range(9, 21):
        a.cell(row=r, column=col).value = None
res, vor = wb["Решение"], wb["Воронка"]
for r in free:
    for col in range(1, 14):
        res.cell(row=r + 1, column=col).value = None   # Решение сдвинуто на строку
    for col in range(1, 11):
        vor.cell(row=r + 9, column=col).value = None   # Воронка: канал из Данные!r -> строка r+9

wb.calculation.fullCalcOnLoad = True
wb.save(DST)
print("saved", DST)

# ---- контрольный пересчёт модели на python ----
tot_r = sum(x[1] for x in DATA); tot_v = sum(x[6] for x in DATA); tot_z = sum(x[5] for x in DATA)
roas_avg = tot_v / tot_r
budget_next = round(tot_r * 1.15, -3)
print(f"\nрасходы {tot_r:,} | выручка {tot_v:,} | заказы {tot_z} | ROAS ср. {roas_avg:.2f} "
      f"| бюджет след. года {budget_next:,.0f}")
rows = []
for name, rasx, pok, per, inter, zak, vyr, marg in DATA:
    roas = vyr / rasx
    romi = (vyr * marg - rasx) / rasx
    mult = min(max(roas / roas_avg, 0.5), 2.0)
    rows.append([name, rasx, roas, romi, mult, rasx * mult, vyr, per, inter, zak, pok, marg])
w = sum(r[5] for r in rows)
print(f"{'канал':24s} {'ROAS':>6s} {'ROMI':>7s} {'множ':>5s} {'бюджет→':>10s} {'Δ':>10s} {'Δ%':>7s}  вердикт")
for r in rows:
    nb = round(r[5] / w * budget_next)
    delta = nb - r[1]
    verdict = ("Убыточен — сократить" if r[2] < 1 / r[11]
               else "Лидер — наращивать" if r[2] >= roas_avg else "Ниже среднего — оптимизировать")
    print(f"{r[0]:24s} {r[2]:6.2f} {r[3]:7.2f} {r[4]:5.2f} {nb:10,.0f} {delta:+10,.0f} "
          f"{nb/r[1]-1:+7.0%}  {verdict}")
print("\nворонка итого: показы→переход "
      f"{sum(x[3] for x in DATA)/sum(x[2] for x in DATA):.2%}, "
      f"переход→интерес {sum(x[4] for x in DATA)/sum(x[3] for x in DATA):.2%}, "
      f"интерес→заказ {sum(x[5] for x in DATA)/sum(x[4] for x in DATA):.2%}, "
      f"показ→заказ {sum(x[5] for x in DATA)/sum(x[2] for x in DATA):.3%}")
