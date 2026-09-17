# -*- coding: utf-8 -*-
"""Шаблон Адамчика «Кейс 2. Анализ каналов продаж», заполненный данными проекта.
Срез за месяц. Ozon и WB исключены — продаж на этих площадках не планируется."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

CH       = ["сайт", "директ", "LaModa", "Блогер", "канал n"]
pokazy   = [28000, 130000, 62000, 85000, 96000]
perehody = [2100,  1100,   2480,  1700,  768]
kliki    = [210,   77,     223,   136,   31]
prodazhi = [40,    12,     28,    20,    3]
zatr     = [17000, 24000,  18000, 35000, 22000]

# юнит-экономика из Кейса 1: цена 3390 ₽, себестоимость 851 ₽; у блогера промокод −10%
cena   = [3390, 3390, 3390, 3051, 3390]
sebes  = 851
komiss = [368, 368, 1337, 361, 368]          # комиссия площадки + логистика + эквайринг
marzha = [c - sebes - k for c, k in zip(cena, komiss)]
pribyl = [m * p for m, p in zip(marzha, prodazhi)]
vyr    = [c * p for c, p in zip(cena, prodazhi)]

N     = len(CH)
LAST  = get_column_letter(1 + N)          # последняя колонка с каналом
TOTC  = 2 + N                             # колонка ИТОГО
TOTL  = get_column_letter(TOTC)

wb = Workbook(); ws = wb.active; ws.title = "Воронка"
thin = Side(style="thin", color="B7B7B7"); bd = Border(thin, thin, thin, thin)
hdr, fin, fout = (PatternFill("solid", fgColor=x) for x in ("D9E2F3", "FFF2CC", "E2EFDA"))

ws["A1"] = "АНАЛИЗ КАНАЛОВ ПРОДАЖ"; ws["A1"].font = Font(bold=True, size=14)
ws["B3"] = "Каналы продаж"; ws.merge_cells(f"B3:{LAST}3")
ws["B3"].alignment = Alignment(horizontal="center"); ws["B3"].font = Font(bold=True)

ws["A4"] = "Данные"
for i, name in enumerate(CH):
    ws.cell(row=4, column=2 + i, value=name)
ws.cell(row=4, column=TOTC, value="ИТОГО")
for c in range(1, TOTC + 1):
    cell = ws.cell(row=4, column=c)
    cell.font, cell.fill = Font(bold=True), hdr
    cell.alignment = Alignment(horizontal="center")

RUB, RUB2 = '#,##0\\ "р."', '#,##0.00\\ "р."'
rows_in = [("кол-во показов", pokazy, "#,##0"),
           ("кол-во переходов", perehody, "#,##0"),
           ("кол-во кликов (добавлений в корзину)", kliki, "#,##0"),
           ("кол-во продаж", prodazhi, "#,##0"),
           ("прибыль с продаж", pribyl, RUB),
           ("затрачено", zatr, RUB)]
for r, (label, vals, fmt) in enumerate(rows_in, start=5):
    ws.cell(row=r, column=1, value=label).font = Font(bold=True)
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + i, value=v); c.number_format, c.fill = fmt, fin
    t = ws.cell(row=r, column=TOTC, value=f"=SUM(B{r}:{LAST}{r})")
    t.number_format, t.font = fmt, Font(bold=True)

rows_calc = [("цена перехода",      "={c}10/{c}6", RUB2),
             ("цена клиента",       "={c}10/{c}8", RUB),
             ("Деньги в кассе",     "={c}9-{c}10", RUB),
             ("ROAS на 1 рубль",    "={c}9/{c}10", "0.00"),
             ("показы - переходы",  "={c}6/{c}5",  "0.00%"),
             ("переходы - клики",   "={c}7/{c}6",  "0.00%"),
             ("клики - продажи",    "={c}8/{c}7",  "0.00%"),
             ("переходы - продажи", "={c}8/{c}6",  "0.00%")]
for r, (label, tpl, fmt) in enumerate(rows_calc, start=11):
    ws.cell(row=r, column=1, value=label).font = Font(bold=True)
    for col in range(2, TOTC + 1):
        c = ws.cell(row=r, column=col, value=tpl.format(c=get_column_letter(col)))
        c.number_format, c.fill = fmt, fout
        if col == TOTC:
            c.font = Font(bold=True)

ws["A20"] = "СПРАВОЧНО — как получена «прибыль с продаж» (маржинальная прибыль)"
ws["A20"].font = Font(bold=True, size=12)
ref = [("Цена продажи, ₽", cena, RUB, False),
       ("Себестоимость юнита, ₽ (футболка 780 + принт 46 + упаковка 25)", [sebes]*N, RUB, False),
       ("Комиссия канала + логистика + эквайринг на 1 заказ, ₽", komiss, RUB, False),
       ("Маржа с 1 заказа, ₽", marzha, RUB, False),
       ("Выручка, ₽", vyr, RUB, True)]
for r, (label, vals, fmt, total) in enumerate(ref, start=21):
    ws.cell(row=r, column=1, value=label)
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + i, value=v); c.number_format = fmt
    if total:
        t = ws.cell(row=r, column=TOTC, value=f"=SUM(B{r}:{LAST}{r})")
        t.number_format, t.font = fmt, Font(bold=True)
ws.cell(row=26, column=1, value="ДРР (доля рекламных расходов в выручке), %")
for col in range(2, TOTC + 1):
    L = get_column_letter(col)
    c = ws.cell(row=26, column=col, value=f"={L}10/{L}25"); c.number_format = "0.0%"

ws.column_dimensions["A"].width = 58
for col in range(2, TOTC + 1):
    ws.column_dimensions[get_column_letter(col)].width = 14
for r in range(4, 19):
    for col in range(1, TOTC + 1):
        ws.cell(row=r, column=col).border = bd
wb.calculation.fullCalcOnLoad = True
wb.save("Кейс2_Анализ_каналов_продаж_заполнено.xlsx")

# CSV с формулами — для загрузки в Google Таблицы
grid = [[""] * TOTC for _ in range(26)]
grid[0][0] = "АНАЛИЗ КАНАЛОВ ПРОДАЖ"
grid[2][1] = "Каналы продаж"
grid[3] = ["Данные"] + CH + ["ИТОГО"]
for i, (label, vals, _) in enumerate(rows_in):
    r = 5 + i
    grid[r-1] = [label] + [str(v) for v in vals] + [f"=SUM(B{r}:{LAST}{r})"]
for i, (label, tpl, _) in enumerate(rows_calc):
    grid[10+i] = [label] + [tpl.format(c=get_column_letter(col)) for col in range(2, TOTC + 1)]
grid[19][0] = "СПРАВОЧНО — как получена «прибыль с продаж» (маржинальная прибыль)"
for i, (label, vals, _, total) in enumerate(ref):
    r = 21 + i
    grid[r-1] = [label] + [str(v) for v in vals] + ([f"=SUM(B{r}:{LAST}{r})"] if total else [""])
grid[25] = ["ДРР (доля рекламных расходов в выручке), %"] + \
           [f"={get_column_letter(c)}10/{get_column_letter(c)}25" for c in range(2, TOTC + 1)]
with open("Кейс2_Анализ_каналов_продаж_заполнено.csv", "w", encoding="utf-8", newline="") as f:
    csv.writer(f).writerows(grid)

print(f"каналов {N} | продажи {sum(prodazhi)} | выручка {sum(vyr):,} | прибыль {sum(pribyl):,} "
      f"| затрачено {sum(zatr):,} | касса {sum(pribyl)-sum(zatr):,} | ROAS {sum(pribyl)/sum(zatr):.2f}")
